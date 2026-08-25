"""
DART 공시서류 원문(XML) → 엑셀 추출 스크립트

DART에서 "공시서류 원본파일 다운로드"로 받은 zip(사업보고서 등)을 읽어,
원하는 목차(섹션)의 표들을 엑셀 파일로 추출한다.

사용법:
    pip install openpyxl

    # 1) 보고서 목차 확인 (어떤 섹션이 있는지)
    python dart_report_to_excel.py 사업보고서.zip --toc

    # 2) 특정 섹션의 표를 엑셀로 추출 (섹션명은 일부만 입력해도 됨)
    python dart_report_to_excel.py 사업보고서.zip --section "매출 및 수주상황" --out 수주.xlsx

    # 3) 여러 연도 zip을 한 번에 → 연도별 시트 생성
    python dart_report_to_excel.py *.zip --section "요약재무정보" --out 요약재무.xlsx

    # 4) 섹션을 여러 개 지정
    python dart_report_to_excel.py *.zip --section "요약재무정보" --section "배당" --out out.xlsx

    # 5) 표 사이의 본문 텍스트도 함께 추출
    python dart_report_to_excel.py 사업보고서.zip --section "사업의 개요" --with-text --out out.xlsx
"""

import argparse
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# 표의 셀로 취급하는 태그 (TD=일반, TH=헤더, TU=단위/자동입력, TE=기타)
CELL_TAGS = {"TD", "TH", "TU", "TE"}
ROW_TAG = "TR"
SECTION_RE = re.compile(r"^SECTION-\d+$")

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
TITLE_FONT = Font(bold=True, size=12)
CAPTION_FONT = Font(bold=True, color="1F4E79")


# ──────────────────────────────────────────────────────────────
# XML 로딩
# ──────────────────────────────────────────────────────────────

def read_xml_bytes(raw: bytes) -> str:
    """DART 원문은 utf-8 또는 euc-kr. 인코딩 선언/디코딩을 자동 판별."""
    head = raw[:200].decode("ascii", errors="ignore").lower()
    encodings = []
    m = re.search(r'encoding="([\w-]+)"', head)
    if m:
        encodings.append(m.group(1))
    encodings += ["utf-8", "cp949", "euc-kr"]
    for enc in encodings:
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="replace")


def parse_document(raw: bytes) -> ET.Element:
    text = read_xml_bytes(raw)
    # 인코딩 선언 제거 후 문자열로 파싱 (선언과 실제 인코딩 불일치 대비)
    text = re.sub(r"<\?xml[^?]*\?>", "", text, count=1)
    # DART 원문은 엄밀한 XML이 아니라서 정리가 필요:
    #  1) &cr; 같은 비표준 엔티티 → &amp;cr; (줄바꿈 표기)
    text = re.sub(r"&(?!lt;|gt;|amp;|quot;|apos;|#)", "&amp;", text)
    #  2) 본문 속 이스케이프 안 된 '<' (예: "< 로드맵 >", "<당사 포트폴리오>")
    #     실제 태그는 <영문자, </영문자, <!, <? 로만 시작한다.
    text = re.sub(r"<(?![A-Za-z!?]|/[A-Za-z])", "&lt;", text)
    try:
        return ET.fromstring(text)
    except ET.ParseError:
        # 3) 그래도 실패하면: 실제 쓰이는 태그명만 허용하고 나머지 <...>는 이스케이프
        allowed = r"[A-Z][A-Z0-9-]*"
        text = re.sub(
            rf"<(?!/?(?:{allowed})(?:\s|/?>))",
            "&lt;", text)
        return ET.fromstring(text)


def expand_inputs(paths):
    """폴더가 섞여 있으면 폴더 안의 zip/xml 파일로 펼친다."""
    out = []
    for p in paths:
        p = Path(p)
        if p.is_dir():
            out += sorted(p.glob("*.zip")) + sorted(p.glob("*.xml"))
        else:
            out.append(p)
    return out


def load_reports(paths, log=print):
    """zip/xml/폴더 경로들 → [(라벨, 루트 Element)] 목록. zip 안에서는 본문 XML만 사용
    (파일명에 _00760 같은 접미사가 붙은 것은 감사보고서 등 첨부문서)."""
    reports = []
    for p in expand_inputs(paths):
        if not p.exists():
            log(f"[경고] 파일 없음: {p}")
            continue
        if p.suffix.lower() == ".zip":
            with zipfile.ZipFile(p) as zf:
                xml_names = [n for n in zf.namelist() if n.lower().endswith(".xml")]
                mains = [n for n in xml_names if "_" not in Path(n).stem]
                target = mains[0] if mains else min(xml_names, key=len)
                raw = zf.read(target)
        else:
            raw = p.read_bytes()
        try:
            root = parse_document(raw)
        except ET.ParseError as e:
            log(f"[경고] XML 파싱 실패({p.name}): {e}")
            continue
        reports.append((report_label(root, p.name), root))
    # 연도순 정렬
    reports.sort(key=lambda r: r[0])
    return reports


def report_label(root: ET.Element, filename: str) -> str:
    """시트 이름에 쓸 라벨: 사업연도(예: 2025) 우선, 없으면 파일명에서 추출."""
    for tu in root.iter("TU"):
        if tu.get("AUNIT") == "PERIODTO" and tu.get("AUNITVALUE"):
            return tu.get("AUNITVALUE")[:4]
    m = re.search(r"(20\d{2})[._]", filename)
    return m.group(1) if m else Path(filename).stem[:8]


# ──────────────────────────────────────────────────────────────
# 목차/섹션 탐색
# ──────────────────────────────────────────────────────────────

def build_parent_map(root: ET.Element) -> dict:
    return {child: parent for parent in root.iter() for child in parent}

def iter_toc(root: ET.Element):
    """(제목, SECTION 요소) 목록. TITLE ATOC='Y' 가 목차 항목."""
    parents = build_parent_map(root)
    for title in root.iter("TITLE"):
        if title.get("ATOC") != "Y":
            continue
        text = clean_text("".join(title.itertext()))
        section = parents.get(title)
        while section is not None and not SECTION_RE.match(section.tag):
            section = parents.get(section)
        if section is not None and text:
            yield text, section


def find_sections(root: ET.Element, keyword: str):
    """목차 제목에 keyword가 포함된 섹션들 반환 (공백 무시 비교)."""
    key = re.sub(r"\s+", "", keyword)
    hits = []
    for text, section in iter_toc(root):
        if key in re.sub(r"\s+", "", text):
            hits.append((text, section))
    return hits


# ──────────────────────────────────────────────────────────────
# 표 → 2차원 그리드
# ──────────────────────────────────────────────────────────────

def clean_text(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def cell_value(text: str):
    """'1,234' → 1234, '(567)' → -567 등 숫자로 변환 가능하면 숫자로."""
    t = text.strip()
    if not t or t == "-":
        return t
    neg = False
    if t.startswith("(") and t.endswith(")"):
        neg, t = True, t[1:-1].strip()
    if t.startswith(("△", "-")):
        neg, t = True, t[1:].strip()
    num = t.replace(",", "")
    if re.fullmatch(r"\d+(\.\d+)?", num):
        val = float(num) if "." in num else int(num)
        return -val if neg else val
    return text.strip()


def table_to_grid(table: ET.Element):
    """<TABLE> → (2차원 리스트, 헤더 행 수). ROWSPAN/COLSPAN 반영."""
    grid = {}          # (row, col) -> value
    occupied = set()   # rowspan/colspan 으로 채워진 자리
    header_rows = 0
    r = 0
    for part in table:
        if part.tag not in ("THEAD", "TBODY", "TFOOT") and part.tag != ROW_TAG:
            continue
        rows = [part] if part.tag == ROW_TAG else list(part)
        for tr in rows:
            if tr.tag != ROW_TAG:
                continue
            c = 0
            for cell in tr:
                if cell.tag not in CELL_TAGS:
                    continue
                while (r, c) in occupied:
                    c += 1
                text = clean_text("".join(cell.itertext()))
                rowspan = int(cell.get("ROWSPAN", 1) or 1)
                colspan = int(cell.get("COLSPAN", 1) or 1)
                grid[(r, c)] = cell_value(text)
                for dr in range(rowspan):
                    for dc in range(colspan):
                        occupied.add((r + dr, c + dc))
            if part.tag == "THEAD":
                header_rows = r + 1
            r += 1
    if not grid:
        return [], 0
    n_rows = max(k[0] for k in grid) + 1
    n_cols = max(k[1] for k in grid) + 1
    out = [[grid.get((i, j), "") for j in range(n_cols)] for i in range(n_rows)]
    return out, header_rows


# ──────────────────────────────────────────────────────────────
# 섹션 → 시트
# ──────────────────────────────────────────────────────────────

def iter_section_content(section: ET.Element, with_text: bool):
    """섹션 안의 내용물을 문서 순서대로 (종류, 내용) 으로 반환.
    종류: 'table' | 'caption' | 'text'. 표 내부로는 내려가지 않아 중복을 막는다."""
    def walk(el):
        for child in el:
            if child.tag == "TABLE":
                yield "table", child
            elif child.tag == "TABLE-GROUP":
                cap = child.get("ACAPTION") or ""
                if cap:
                    yield "caption", cap
                yield from walk(child)
            elif child.tag == "P":
                if with_text:
                    text = clean_text("".join(child.itertext()))
                    if text:
                        yield "text", text
            elif child.tag == "TITLE":
                continue
            else:
                yield from walk(child)
    yield from walk(section)


def sanitize_sheet_name(name: str) -> str:
    name = re.sub(r'[\\/*?:\[\]]', "", name)
    return name[:31] or "Sheet"


def write_section_sheet(wb: Workbook, sheet_name: str, section_title: str,
                        section: ET.Element, with_text: bool):
    ws = wb.create_sheet(sanitize_sheet_name(sheet_name))
    ws.cell(row=1, column=1, value=section_title).font = TITLE_FONT
    row = 3
    max_col = 1
    seen_tables = set()
    for kind, content in iter_section_content(section, with_text):
        if kind == "caption":
            ws.cell(row=row, column=1, value=content).font = CAPTION_FONT
            row += 1
        elif kind == "text":
            ws.cell(row=row, column=1, value=content)
            row += 1
        elif kind == "table":
            if id(content) in seen_tables:
                continue
            seen_tables.add(id(content))
            grid, header_rows = table_to_grid(content)
            if not grid:
                continue
            for i, line in enumerate(grid):
                for j, val in enumerate(line):
                    c = ws.cell(row=row + i, column=1 + j, value=val)
                    if i < header_rows:
                        c.font = Font(bold=True)
                        c.fill = HEADER_FILL
                        c.alignment = Alignment(horizontal="center")
                max_col = max(max_col, len(line))
            row += len(grid) + 1  # 표 사이 한 줄 띄움
    # 대략적인 열 너비
    for col in range(1, max_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    return ws


# ──────────────────────────────────────────────────────────────
# 추출 (CLI/GUI 공용)
# ──────────────────────────────────────────────────────────────

def toc_lines(reports):
    """보고서별 목차를 문자열 리스트로 반환."""
    lines = []
    for label, root in reports:
        name = clean_text("".join(next(root.iter("COMPANY-NAME"), ET.Element("x")).itertext()))
        doc = clean_text("".join(next(root.iter("DOCUMENT-NAME"), ET.Element("x")).itertext()))
        lines.append(f"===== [{label}] {name} {doc} =====")
        for text, _ in iter_toc(root):
            lines.append("  " + text)
        lines.append("")
    return lines


def extract_to_workbook(reports, keywords, with_text=False, log=print):
    """섹션 키워드들로 표를 추출해 (Workbook, 시트 수) 반환."""
    wb = Workbook()
    wb.remove(wb.active)
    n_sheets = 0
    for label, root in reports:
        for keyword in keywords:
            hits = find_sections(root, keyword)
            if not hits:
                log(f"[{label}] '{keyword}' 섹션 없음 — 건너뜀")
                continue
            for title, section in hits:
                short = re.sub(r"^[IVX0-9\.\s]+", "", title)  # 앞의 번호 제거
                sheet = f"{label}_{short}"
                base = sheet
                k = 2
                while sanitize_sheet_name(sheet) in wb.sheetnames:
                    sheet = f"{base}_{k}"
                    k += 1
                write_section_sheet(wb, sheet, f"[{label}] {title}", section, with_text)
                n_sheets += 1
                log(f"[{label}] '{title}' → 시트 저장")
    return wb, n_sheets


# ──────────────────────────────────────────────────────────────
# main
# ──────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(description="DART 원문 zip/xml에서 섹션 표를 엑셀로 추출")
    ap.add_argument("files", nargs="+", help="DART 원문 zip/xml 파일 또는 폴더 (폴더면 안의 zip 전부)")
    ap.add_argument("--toc", action="store_true", help="목차만 출력하고 종료")
    ap.add_argument("--section", action="append", default=[],
                    help="추출할 섹션명 키워드 (여러 번 지정 가능, 일부 문자열이면 됨)")
    ap.add_argument("--with-text", action="store_true", help="표 외의 본문 문단도 포함")
    ap.add_argument("--out", default="dart_extract.xlsx", help="출력 엑셀 파일명")
    args = ap.parse_args()

    reports = load_reports(args.files)
    if not reports:
        sys.exit("읽을 수 있는 보고서가 없습니다.")

    if args.toc:
        print("\n".join(toc_lines(reports)))
        return

    if not args.section:
        sys.exit("--section 키워드를 지정하세요. (--toc 로 목차를 먼저 확인)")

    wb, n_sheets = extract_to_workbook(reports, args.section, args.with_text)
    if n_sheets == 0:
        sys.exit("추출된 섹션이 없습니다. --toc 로 정확한 섹션명을 확인하세요.")
    wb.save(args.out)
    print(f"\n완료: {args.out} ({n_sheets}개 시트)")


if __name__ == "__main__":
    main()
