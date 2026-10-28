"""
DART 수주공시(단일판매ㆍ공급계약체결) 수집·정리 → 엑셀

대상(기본): HD현대중공업, HD현대미포 (미포는 2025-12-01 HD현대중공업에 합병되어 이후 공시 없음)
옵션으로 HD한국조선해양(자회사 HD현대삼호 수주공시 포함)도 추가 가능.

사용법:
    pip install requests openpyxl

    # 1) DART에서 공시 원문을 받아 정리 (원문 XML은 --raw-dir 에 저장)
    python ship_orders_dart.py --api-key YOUR_DART_KEY --years 6 \
        --raw-dir dart_data/수주공시 --out 결과물/수주공시_현대중공업_미포.xlsx

    # 2) 이미 받아둔 원문 XML만으로 다시 정리 (인터넷/API 키 불필요)
    python ship_orders_dart.py --from-dir dart_data/수주공시 --out 결과물/수주공시_현대중공업_미포.xlsx

    # 한국조선해양(삼호 포함)까지
    python ship_orders_dart.py --api-key KEY --corps 현대중공업 미포 한국조선해양 ...

출력 엑셀 시트:
    공시목록   : 공시별 1행 (정정공시는 최신 정정본만 남김, 해지·변경 공시는 표시)
    분기집계   : 회사 × 계약분기 별 건수·척수·계약금액(십억원)
    선종집계   : 회사 × 선종 별 건수·척수·계약금액·척당금액
"""

import argparse
import html
import io
import re
import time
import zipfile
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

CORPS = {
    "현대중공업":   ("01390344", "HD현대중공업"),
    "미포":         ("00164609", "HD현대미포"),
    "한국조선해양": ("00164830", "HD한국조선해양"),
}
DART_BASE = "https://opendart.fss.or.kr/api"

VESSEL_KEYWORDS = [
    ("LNG BV", "LNG BV선"), ("벙커링", "LNG BV선"),
    ("LNG", "LNG선"), ("액화천연가스", "LNG선"),
    ("VLAC", "VLAC"), ("초대형 암모니아", "VLAC"), ("초대형암모니아", "VLAC"),
    ("VLGC", "LPG선"), ("LPG", "LPG선"), ("액화석유가스", "LPG선"), ("암모니아", "LPG선"),
    ("에탄", "가스선"), ("LCO2", "LCO2선"), ("이산화탄소", "LCO2선"),
    ("컨테이너", "컨테이너선"),
    ("VLCC", "VLCC"), ("초대형 원유", "VLCC"), ("초대형원유", "VLCC"),
    ("셔틀탱커", "탱커"), ("원유운반", "탱커"), ("원유 운반", "탱커"), ("수에즈막스", "탱커"),
    ("아프라막스", "탱커"), ("석유화학제품", "P/C선"), ("PC선", "P/C선"), ("P/C", "P/C선"),
    ("MR", "P/C선"), ("제품운반", "P/C선"), ("탱커", "탱커"),
    ("자동차운반", "PCTC"), ("PCTC", "PCTC"),
    ("벌크", "벌크선"), ("살물", "벌크선"),
    ("FPSO", "해양"), ("FLNG", "해양"), ("해양플랜트", "해양"), ("플랫폼", "해양"), ("풍력", "해양"),
    ("잠수함", "특수선"), ("호위함", "특수선"), ("구축함", "특수선"), ("함정", "특수선"),
    ("군수지원함", "특수선"), ("경비함", "특수선"), ("쇄빙", "특수선"),
    ("엔진", "엔진"), ("발전", "엔진"),
]


# ── DART API ──────────────────────────────────────────────────────────────

def get_list(api_key, corp_code, bgn_de, end_de):
    import requests
    out, page = [], 1
    while True:
        try:
            r = requests.get(f"{DART_BASE}/list.json", params={
                "crtfc_key": api_key, "corp_code": corp_code,
                "bgn_de": bgn_de, "end_de": end_de,
                "page_count": 100, "page_no": page,
            }, timeout=20).json()
        except Exception as e:
            print(f"    [오류] 목록 조회: {e}")
            break
        if r.get("status") != "000":
            if r.get("status") != "013":  # 013 = 조회된 데이터 없음
                print(f"    [DART] {r.get('status')} {r.get('message')}")
            break
        for item in r.get("list", []):
            nm = item.get("report_nm", "")
            if "단일판매" in nm or "공급계약" in nm:
                out.append(item)
        if page * 100 >= int(r.get("total_count", 0)):
            break
        page += 1
        time.sleep(0.3)
    return out


def get_document(api_key, rcept_no):
    import requests
    try:
        r = requests.get(f"{DART_BASE}/document.xml",
                         params={"crtfc_key": api_key, "rcept_no": rcept_no}, timeout=30)
    except Exception as e:
        print(f"    [오류] 원문 다운로드: {e}")
        return None
    if r.content[:2] == b"PK":
        with zipfile.ZipFile(io.BytesIO(r.content)) as z:
            names = sorted((n for n in z.namelist() if n.endswith(".xml")), key=len)
            if names:
                return z.read(names[0])
    return None


# ── 원문 → 셀 토큰 ────────────────────────────────────────────────────────

def decode(raw: bytes) -> str:
    head = raw[:200].decode("ascii", errors="ignore").lower()
    m = re.search(r'encoding="([\w-]+)"', head)
    for enc in ([m.group(1)] if m else []) + ["utf-8", "cp949", "euc-kr"]:
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="replace")


def tokens(text: str):
    """표 셀(TD/TH/TE/TU) 단위로 텍스트 토큰을 순서대로 반환."""
    t = text.replace("&cr;", " ")
    t = re.sub(r"<(TD|TH|TE|TU)\b[^>]*>", "\x01", t)
    t = re.sub(r"</(TD|TH|TE|TU|TR|P|TITLE)>", "\x01", t)
    t = re.sub(r"<[^>]+>", "", t)
    t = html.unescape(t)
    out = []
    for s in t.split("\x01"):
        s = re.sub(r"\s+", " ", s.replace("\xa0", " ")).strip()
        if s:
            out.append(s)
    return out


def norm(s):
    s = re.sub(r"^\s*(\d+|[가-하]|[①-⑳])\s*[\.\)]\s*", "", s)  # "1. ", "가. " 같은 항목번호 제거
    return re.sub(r"[\sㆍ·\.\-]", "", s)


def after(tok, label, *, last=True, skip=0, pred=None):
    """label 셀 다음에 오는 값. 정정공시는 정정 후 본문이 뒤에 오므로 기본은 마지막 매치."""
    key = norm(label)
    hits = [i for i, s in enumerate(tok) if norm(s).startswith(key)]
    for i in (reversed(hits) if last else hits):
        vals = tok[i + 1:i + 8]
        vals = [v for v in vals if v not in ("-",)]
        if pred:
            vals = [v for v in vals if pred(v)]
        if len(vals) > skip:
            return vals[skip]
    return ""


def is_num(v):
    return bool(re.fullmatch(r"\(?-?[\d,]+(\.\d+)?\)?", v.replace(" ", "")))


def to_num(v):
    v = v.replace(",", "").replace(" ", "")
    neg = v.startswith("(") or v.startswith("-")
    v = v.strip("()-")
    try:
        x = float(v)
    except ValueError:
        return None
    return -x if neg else x


def to_date(v):
    m = re.search(r"(\d{4})\s*[.\-년/]\s*(\d{1,2})\s*[.\-월/]\s*(\d{1,2})", v)
    if m:
        try:
            return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    return None


def vessel_type(text):
    up = text.upper()
    for kw, label in VESSEL_KEYWORDS:
        if kw.upper() in up:
            return label
    return ""


def ship_count(text):
    m = re.search(r"(\d+)\s*척", text)
    return int(m.group(1)) if m else None


def usd_amount(text):
    """본문에서 달러 계약금액 추정 (백만달러)."""
    for pat, mul in [(r"(?:USD|US\$|\$)\s*([\d,\.]+)\s*(?:백만|M\b|million)", 1),
                     (r"([\d,\.]+)\s*백만\s*(?:달러|USD|불)", 1),
                     (r"([\d,\.]+)\s*억\s*(?:달러|불)", 100),
                     (r"(?:USD|US\$)\s*([\d,]{7,})", 1e-6)]:
        m = re.search(pat, text, re.I)
        if m:
            v = to_num(m.group(1))
            if v:
                return round(v * mul, 2)
    return None


def parse(raw: bytes, meta: dict):
    text = decode(raw)
    tok = tokens(text)
    flat = " ".join(tok)
    d = dict(meta)

    content = after(tok, "판매ㆍ공급계약 내용") or after(tok, "판매ㆍ공급계약 구분", skip=0) \
        or after(tok, "계약내용")
    detail = after(tok, "세부내용") or ""
    d["계약내용"] = (content + (" / " + detail if detail and detail != content else "")).strip(" /")

    amt = after(tok, "계약금액", pred=is_num)
    d["계약금액(원)"] = to_num(amt) if amt else None
    sales = after(tok, "최근매출액", pred=is_num)
    d["최근매출액(원)"] = to_num(sales) if sales else None
    ratio = after(tok, "매출액대비", pred=is_num)
    d["매출액대비(%)"] = to_num(ratio) if ratio else None

    d["계약상대방"] = after(tok, "계약상대방") or after(tok, "계약상대")
    d["판매ㆍ공급지역"] = after(tok, "판매ㆍ공급지역")
    d["계약시작일"] = to_date(after(tok, "시작일"))
    d["계약종료일"] = to_date(after(tok, "종료일"))
    d["계약(수주)일자"] = to_date(after(tok, "계약(수주)일자") or after(tok, "계약일자"))

    basis = " ".join([d["계약내용"], d.get("공시제목", "")])
    d["선종"] = vessel_type(basis) or vessel_type(flat[:4000])
    d["척수"] = ship_count(d["계약내용"]) or ship_count(flat)
    d["계약금액(백만달러,본문)"] = usd_amount(flat)
    d["구분"] = ("해지" if "해지" in d.get("공시제목", "") else
                 "정정" if "정정" in d.get("공시제목", "") else "신규")
    return d


# ── 수집 / 로드 ────────────────────────────────────────────────────────────

def collect(api_key, corp_keys, years, raw_dir):
    today = datetime.today()
    start = today.replace(year=today.year - years, month=1, day=1)
    ranges, cur = [], start
    while cur < today:
        nxt = min(cur + timedelta(days=364), today)
        ranges.append((cur.strftime("%Y%m%d"), nxt.strftime("%Y%m%d")))
        cur = nxt + timedelta(days=1)
    for key in corp_keys:
        code, name = CORPS[key]
        print(f"\n[{name}] {start:%Y-%m-%d} ~ {today:%Y-%m-%d}")
        for bgn, end in ranges:
            items = get_list(api_key, code, bgn, end)
            print(f"  {bgn}~{end}: 수주공시 {len(items)}건")
            for it in items:
                f = Path(raw_dir) / name / f"{it['rcept_dt']}_{it['rcept_no']}.xml"
                if f.exists():
                    continue
                raw = get_document(api_key, it["rcept_no"])
                if raw is None:
                    print(f"    [건너뜀] {it['report_nm']} {it['rcept_no']}")
                    continue
                f.parent.mkdir(parents=True, exist_ok=True)
                f.write_bytes(raw)
                (f.with_suffix(".title.txt")).write_text(it["report_nm"], encoding="utf-8")
                time.sleep(0.5)


def load_dir(raw_dir, corp_keys):
    names = {CORPS[k][1] for k in corp_keys}
    rows = []
    for f in sorted(Path(raw_dir).glob("*/*.xml")):
        corp = f.parent.name
        if corp not in names:
            continue
        rcept_dt, rcept_no = f.stem.split("_", 1)
        tf = f.with_suffix(".title.txt")
        title = tf.read_text(encoding="utf-8") if tf.exists() else ""
        meta = {"회사": corp, "접수일": datetime.strptime(rcept_dt, "%Y%m%d"),
                "접수번호": rcept_no, "공시제목": title}
        try:
            rows.append(parse(f.read_bytes(), meta))
        except Exception as e:
            print(f"  [파싱 실패] {f.name}: {e}")
    return rows


def dedupe(rows):
    """같은 계약의 원공시·정정공시 중 최신(접수번호 최대)만 남김.
    계약 식별: 회사 + 계약(수주)일자 + 계약상대방 (+ 일자 없으면 접수번호)."""
    best = {}
    for r in rows:
        k = (r["회사"], r["계약(수주)일자"] or r["접수번호"], r["계약상대방"])
        if k not in best or r["접수번호"] > best[k]["접수번호"]:
            best[k] = r
    out = sorted(best.values(), key=lambda r: (r["회사"], r["계약(수주)일자"] or r["접수일"]))
    for r in out:
        r["정정반영"] = "O" if r["구분"] == "정정" else ""
    return out


# ── 엑셀 ──────────────────────────────────────────────────────────────────

def quarter(d):
    return f"{d.year}.{(d.month - 1) // 3 + 1}Q" if d else ""


def write_excel(rows, out):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter as gl

    hdr_font = Font(name="Arial", bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    base = Font(name="Arial")

    wb = Workbook()
    ws = wb.active
    ws.title = "공시목록"
    cols = ["회사", "접수일", "접수번호", "공시제목", "구분", "계약(수주)일자", "계약분기", "선종", "척수",
            "계약내용", "계약상대방", "판매ㆍ공급지역", "계약금액(원)", "계약금액(십억원)",
            "척당금액(십억원)", "계약금액(백만달러,본문)", "최근매출액(원)", "매출액대비(%)",
            "계약시작일", "계약종료일(인도)"]
    ws.append(cols)
    for i, r in enumerate(rows, start=2):
        ws.append([r["회사"], r["접수일"], r["접수번호"], r["공시제목"], r["구분"],
                   r["계약(수주)일자"], None, r["선종"], r["척수"], r["계약내용"], r["계약상대방"],
                   r["판매ㆍ공급지역"], r["계약금액(원)"], None, None, r["계약금액(백만달러,본문)"],
                   r["최근매출액(원)"], r["매출액대비(%)"], r["계약시작일"], r["계약종료일"]])
        ws.cell(i, 7).value = f'=IF(F{i}="","",YEAR(F{i})&"."&ROUNDUP(MONTH(F{i})/3,0)&"Q")'
        ws.cell(i, 14).value = f'=IF(M{i}="","",M{i}/1000000000)'
        ws.cell(i, 15).value = f'=IF(OR(N{i}="",I{i}=""),"",N{i}/I{i})'
    last = len(rows) + 1
    fmt = {2: "yyyy-mm-dd", 6: "yyyy-mm-dd", 19: "yyyy-mm-dd", 20: "yyyy-mm-dd",
           13: "#,##0", 14: "#,##0.0", 15: "#,##0.0", 16: "#,##0.0", 17: "#,##0", 18: "0.0"}
    for row in ws.iter_rows(min_row=1, max_row=last):
        for c in row:
            c.font = base
            if c.row == 1:
                c.font, c.fill = hdr_font, hdr_fill
                c.alignment = Alignment(horizontal="center", wrap_text=True)
            elif c.column in fmt:
                c.number_format = fmt[c.column]
    widths = [12, 11, 15, 34, 6, 12, 9, 11, 6, 40, 24, 14, 18, 12, 12, 14, 18, 10, 11, 12]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[gl(j)].width = w
    ws.freeze_panes = "B2"

    # 집계 시트: 공시목록을 참조하는 수식 (해지 공시 제외)
    corps = sorted({r["회사"] for r in rows})
    qs = sorted({quarter(r["계약(수주)일자"]) for r in rows if r["계약(수주)일자"]},
                key=lambda s: (int(s[:4]), s[5]))
    rng = lambda col: f"공시목록!${col}$2:${col}${last}"

    wq = wb.create_sheet("분기집계")
    wq.append(["계약분기"] + [f"{c} {k}" for c in corps for k in ("건수", "척수", "금액(십억원)")])
    for i, q in enumerate(qs, start=2):
        wq.cell(i, 1, q)
        for j, c in enumerate(corps):
            col = 2 + j * 3
            crit = f'{rng("A")},"{c}",{rng("G")},$A{i},{rng("E")},"<>해지"'
            wq.cell(i, col, f"=COUNTIFS({crit})")
            wq.cell(i, col + 1, f"=SUMIFS({rng('I')},{crit})")
            wq.cell(i, col + 2, f"=SUMIFS({rng('N')},{crit})").number_format = "#,##0.0"

    wt = wb.create_sheet("선종집계")
    types = sorted({r["선종"] or "미분류" for r in rows})
    wt.append(["회사", "선종", "건수", "척수", "금액(십억원)", "척당금액(십억원)"])
    i = 2
    for c in corps:
        for t in types:
            tcrit = '""' if t == "미분류" else f'"{t}"'
            crit = f'{rng("A")},"{c}",{rng("H")},{tcrit},{rng("E")},"<>해지"'
            wt.append([c, t, f"=COUNTIFS({crit})", f"=SUMIFS({rng('I')},{crit})",
                       f"=SUMIFS({rng('N')},{crit})", f'=IF(D{i}=0,"",E{i}/D{i})'])
            wt.cell(i, 5).number_format = "#,##0.0"
            wt.cell(i, 6).number_format = "#,##0.0"
            i += 1
    for s in (wq, wt):
        for c in s[1]:
            c.font, c.fill = hdr_font, hdr_fill
        s.freeze_panes = "A2"

    Path(out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"\n엑셀 저장: {out}  (공시 {len(rows)}건)")


def main():
    ap = argparse.ArgumentParser(description="DART 수주공시(단일판매ㆍ공급계약) 정리")
    ap.add_argument("--api-key", default="", help="DART OpenAPI 인증키 (--from-dir만 쓸 땐 불필요)")
    ap.add_argument("--corps", nargs="+", default=["현대중공업", "미포"], choices=list(CORPS))
    ap.add_argument("--years", type=int, default=6)
    ap.add_argument("--raw-dir", default="dart_data/수주공시", help="공시 원문 XML 저장 폴더")
    ap.add_argument("--from-dir", default="", help="저장된 원문만으로 정리 (다운로드 생략)")
    ap.add_argument("--out", default="결과물/수주공시_현대중공업_미포.xlsx")
    a = ap.parse_args()

    raw_dir = a.from_dir or a.raw_dir
    if not a.from_dir:
        if not a.api_key:
            ap.error("--api-key 또는 --from-dir 중 하나는 필요합니다")
        collect(a.api_key, a.corps, a.years, raw_dir)
    rows = dedupe(load_dir(raw_dir, a.corps))
    write_excel(rows, a.out)


if __name__ == "__main__":
    main()
