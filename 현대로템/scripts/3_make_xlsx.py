# -*- coding: utf-8 -*-
"""값 직접 기록 방식 (LibreOffice recalc 불필요)."""
import json, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

M = json.load(open("/tmp/claude-0/-home-user-jooho96/2e90a833-c170-58f0-a6f1-86983b335b93/scratchpad/dataset.json", encoding="utf-8"))
LABS = ["2024Q4", "2025Q2", "2025Q3", "2025Q4", "2026Q1"]
FONT = "맑은 고딕"

# ---------- 분류 ----------
COUNTRY=[("폴란드","폴란드"),("시드니","호주"),("QTMP","호주"),("호주","호주"),
 ("타이베이","대만"),("카오슝","대만"),("도원","대만"),("타이중","대만"),("대만","대만"),
 ("필리핀","필리핀"),("MBTA","미국"),("LA 메트로","미국"),("미국","미국"),
 ("싱가포르","싱가포르"),("웰링턴","뉴질랜드"),("뉴질랜드","뉴질랜드"),
 ("에드먼턴","캐나다"),("캐나다","캐나다"),("우즈베키스탄","우즈베키스탄"),
 ("카이로","이집트"),("알렉산드리아","이집트"),("나그하마디","이집트"),("이집트","이집트"),
 ("이스탄불","터키"),("터키","터키"),("알타이","터키"),
 ("모로코","모로코"),("탄자니아","탄자니아"),("아일랜드","아일랜드"),("이란","이란")]
DOMESTIC=["방위사업청","한화","코레일","한국철도공사","SR","현대제철","현대자동차","공항철도","동북선",
 "대전","지티엑스","서부광역메트로","포스코이앤씨","포스코","수도권광역","우진산전"]
DEF_KW=["전차","K1","K-1","알타이","장갑차","대공포","지휘소용","자주","도하","PBL","성과기반","창정비",
 "외주정비","개척전차","차체포탑","군비청","방위사업청","한화에어로"]
ETC_KW=["제철","CDQ","LNG","자가발전","의장 운반","운반설비","GP신공장"]
RAIL_KW=["전동차","철도","고속철도","고속전철","EMU","트램","경전철","광역철도","도시철도","MRT","LRT",
 "메트로","E&M","유지보수","동차","객차","기관차","통근형","2층 전동차","차량운행시스템","광역급행",
 "디젤전동차","디젤동차","신호","LTA","육상교통","J151"]
def country_of(t):
    for kw,c in COUNTRY:
        if kw in t: return c
    return None
def classify(name, party):
    t=name+" "+party; c=country_of(t)
    if any(k in t for k in DEF_KW): dae="방산"
    elif any(k in name for k in ETC_KW) or "현대자동차" in party: dae="기타"
    elif any(k in t for k in RAIL_KW): dae="철도"
    else: dae="철도"
    dom=any(k in party for k in DOMESTIC)
    if c: so,guk="국외",c
    elif dom: so,guk="국내",""
    else: so,guk="국내",""
    return dae,so,guk
def n(x): return x if isinstance(x,(int,float)) else None

# ---------- 레코드 ----------
recs=[]
for key,m in M.items():
    name=sorted(m["names"], key=len)[-1]
    party=m["party"][0] if m["party"] else ""
    dae,so,guk=classify(name,party)
    curr={l:n(m["curr"].get(l)) for l in LABS}
    cum={l:n(m["cum"].get(l)) for l in LABS}
    amt=n(m.get("amt_krw")); amt_raw=m.get("amt_raw","")
    subs=sorted(set(m.get("subnames",[]))) if isinstance(m.get("subnames"),list) else []
    # 분기 계산
    def sub(a,b): return None if a is None else a-(b or 0)
    q={"2024년":curr["2024Q4"], "2025_상반기":curr["2025Q2"],
       "2025_3Q":sub(curr["2025Q3"],curr["2025Q2"]),
       "2025_4Q":sub(curr["2025Q4"],curr["2025Q3"]),
       "2026_1Q":curr["2026Q1"]}
    notes=[]
    if subs: notes.append("통합계약: "+", ".join(s[:22] for s in subs))
    if amt_raw and re.match(r"^(USD|EUR|JPY|GBP|AUD|SGD|TWD|NZD|CAD)", amt_raw):
        notes.append(f"계약총액 외화기준({amt_raw})")
    if curr.get("2025Q3") is not None and curr.get("2025Q2") is None and curr.get("2024Q4") is None:
        notes.append("표편입 2025.3Q(해당값에 이전기간 포함가능)")
    recs.append(dict(dae=dae,so=so,guk=guk,name=name,party=party,start=m.get("start",""),
        end=m.get("end",""),amt=amt,curr=curr,cum=cum,cum_latest=n(m.get("cum_latest")),q=q,note=" / ".join(notes)))

order={"철도":0,"방산":1,"기타":2}
recs.sort(key=lambda r:(order[r["dae"]], 0 if r["so"]=="국내" else 1, -(r["cum_latest"] or 0)))

# ---------- 스타일 ----------
navy=PatternFill("solid",fgColor="1F3864"); blue=PatternFill("solid",fgColor="2E5496")
rail=PatternFill("solid",fgColor="DDEBF7"); defe=PatternFill("solid",fgColor="FCE4D6")
etc=PatternFill("solid",fgColor="E2EFDA"); subf=PatternFill("solid",fgColor="D9E1F2"); grand=PatternFill("solid",fgColor="BDD7EE")
hf=Font(name=FONT,bold=True,color="FFFFFF",size=9); bf=Font(name=FONT,bold=True,size=9)
nf=Font(name=FONT,size=9); sf=Font(name=FONT,size=8,color="808080")
thin=Side(style="thin",color="BFBFBF"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
ctr=Alignment("center","center",wrap_text=True); lft=Alignment("left","center",wrap_text=True); rgt=Alignment("right","center")
WON="#,##0;(#,##0);-"
def fillfor(dae): return rail if dae=="철도" else defe if dae=="방산" else etc

wb=openpyxl.Workbook()

# ===== 메인 =====
ws=wb.active; ws.title="계약별_분기매출인식"
cols=["No","대분류","소분류","국가","계약명","계약상대방","계약시작일","계약종료일","계약총액",
      "최신누적\n매출인식","진행률","2024년\n(연간)","2025\n상반기","2025\n3분기","2025\n4분기","2026\n1분기","비고"]
ncol=len(cols)
ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=ncol)
c=ws.cell(1,1,"현대로템  분기별 단일판매·공급계약  매출인식 정리"); c.font=Font(name=FONT,bold=True,color="FFFFFF",size=14); c.fill=navy; c.alignment=lft; ws.row_dimensions[1].height=26
ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=ncol)
c=ws.cell(2,1,"단위: 백만원(百萬KRW)  ·  출처: 각 정기보고서 「XI-1-나-(1) 단일판매·공급계약 공시 진행」  ·  분기값 = 보고서 '당기(YTD)' 차분  ·  세부 원본은 [원자료_당기누적] 시트 참조")
c.font=sf; c.alignment=lft; ws.row_dimensions[2].height=15
HDR=3
for j,h in enumerate(cols,1):
    cc=ws.cell(HDR,j,h); cc.font=hf; cc.fill=blue if j>=12 else navy; cc.alignment=ctr; cc.border=bd
ws.row_dimensions[HDR].height=30; ws.freeze_panes="E4"

def put(r,j,v,fill=None,fmt=None,al=None,font=nf):
    cc=ws.cell(r,j,v); cc.font=font; cc.border=bd
    cc.alignment=al or (rgt if isinstance(v,(int,float)) else ctr)
    if fill: cc.fill=fill
    if fmt: cc.number_format=fmt
    return cc

QKEYS=["2024년","2025_상반기","2025_3Q","2025_4Q","2026_1Q"]
r=HDR+1; No=0; cur_cat=None; cat_first=None
subtot={}  # dae -> [amt,cum,q...]
def emit_subtotal(dae, upto_row, first_row):
    global r
    agg=[0]*7  # amt,cum,4q? -> amt,cum, q0..q4(5) =7
    agg=[0,0,0,0,0,0,0]
    for rec in recs:
        if rec["dae"]!=dae: continue
        agg[0]+=rec["amt"] or 0; agg[1]+=rec["cum_latest"] or 0
        for k,qk in enumerate(QKEYS): agg[2+k]+=rec["q"][qk] or 0
    lab={"철도":"철도 소계","방산":"방산 소계","기타":"기타 소계"}[dae]
    put(r,5,lab,fill=subf,al=rgt,font=bf)
    for j in range(1,ncol+1): ws.cell(r,j).fill=subf; ws.cell(r,j).border=bd
    put(r,9,agg[0],fill=subf,fmt=WON,font=bf); put(r,10,agg[1],fill=subf,fmt=WON,font=bf)
    for k in range(5): put(r,12+k,agg[2+k],fill=subf,fmt=WON,font=bf)
    r+=1

cats=["철도","방산","기타"]
for ci,dae in enumerate(cats):
    group=[rec for rec in recs if rec["dae"]==dae]
    for rec in group:
        No+=1
        f=fillfor(dae)
        put(r,1,No,al=ctr); put(r,2,rec["dae"],fill=f,al=ctr); put(r,3,rec["so"],fill=f,al=ctr); put(r,4,rec["guk"],fill=f,al=ctr)
        put(r,5,rec["name"],fill=f,al=lft); put(r,6,rec["party"],fill=f,al=lft)
        put(r,7,rec["start"],al=ctr); put(r,8,rec["end"],al=ctr)
        put(r,9,rec["amt"],fmt=WON)
        put(r,10,rec["cum_latest"],fmt=WON)
        prog=(rec["cum_latest"]/rec["amt"]) if (rec["amt"] and rec["cum_latest"] is not None) else None
        put(r,11,prog,fmt="0.0%")
        for k,qk in enumerate(QKEYS): put(r,12+k,rec["q"][qk],fmt=WON)
        put(r,17,rec["note"],al=lft,font=sf)
        r+=1
    emit_subtotal(dae,r-1,None)

# 총합계
put(r,5,"총 합계",fill=grand,al=rgt,font=bf)
for j in range(1,ncol+1): ws.cell(r,j).fill=grand; ws.cell(r,j).border=bd
ta=sum(x["amt"] or 0 for x in recs); tc=sum(x["cum_latest"] or 0 for x in recs)
put(r,9,ta,fill=grand,fmt=WON,font=bf); put(r,10,tc,fill=grand,fmt=WON,font=bf)
for k,qk in enumerate(QKEYS): put(r,12+k,sum(x["q"][qk] or 0 for x in recs),fill=grand,fmt=WON,font=bf)
widths=[4,6,6,9,34,22,11,11,12,12,7,10,9,9,9,9,26]
for j,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(j)].width=w

# ===== 원자료 =====
raw=wb.create_sheet("원자료_당기누적")
rcols=["대분류","소분류","국가","계약명","계약상대방","신고일자","계약시작일","계약종료일","계약총액"]+[f"당기_{l}" for l in LABS]+[f"누적_{l}" for l in LABS]
raw.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(rcols))
c=raw.cell(1,1,"[원자료] 보고서별 당기(YTD)·누적 매출인식 — 각 정기보고서 XI-1-나-(1) 표에서 직접 추출 (단위: 백만원)")
c.font=Font(name=FONT,bold=True,color="FFFFFF",size=11); c.fill=navy; c.alignment=lft; raw.row_dimensions[1].height=20
raw.merge_cells(start_row=2,start_column=10,end_row=2,end_column=14); gc=raw.cell(2,10,"당기 매출인식 (회계연도 누계 YTD)"); gc.font=hf; gc.fill=blue; gc.alignment=ctr
raw.merge_cells(start_row=2,start_column=15,end_row=2,end_column=19); gc=raw.cell(2,15,"누적 매출인식 (계약개시~기말)"); gc.font=hf; gc.fill=blue; gc.alignment=ctr
for j in range(1,10):
    cc=raw.cell(2,j,rcols[j-1]); cc.font=hf; cc.fill=navy; cc.alignment=ctr
labhdr={"2024Q4":"2024사업\n(FY24)","2025Q2":"2025반기\n(H1)","2025Q3":"2025.3Q\n(9M)","2025Q4":"2025사업\n(FY25)","2026Q1":"2026.1Q\n(3M)"}
H2=3
for j,h in enumerate(rcols,1):
    disp=labhdr[h.split("_")[1]] if ("당기_" in h or "누적_" in h) else h
    cc=raw.cell(H2,j,disp); cc.font=hf; cc.fill=blue if j>=10 else navy; cc.alignment=ctr; cc.border=bd
raw.row_dimensions[H2].height=26
rr=H2+1
for rec in recs:
    m=next(mm for kk,mm in M.items() if sorted(mm["names"],key=len)[-1]==rec["name"] and (mm.get("start","")==rec["start"]))
    vals=[rec["dae"],rec["so"],rec["guk"],rec["name"],rec["party"],m.get("sinbo",""),rec["start"],rec["end"],rec["amt"]]
    for j,v in enumerate(vals,1):
        cc=raw.cell(rr,j,v); cc.font=nf; cc.border=bd
        cc.alignment=lft if j in (4,5) else (rgt if j==9 else ctr)
        if j==9: cc.number_format=WON
        if j<=5: cc.fill=fillfor(rec["dae"])
    for k,l in enumerate(LABS):
        cc=raw.cell(rr,10+k,rec["curr"][l]); cc.font=nf; cc.border=bd; cc.alignment=rgt; cc.number_format=WON
    for k,l in enumerate(LABS):
        cc=raw.cell(rr,15+k,rec["cum"][l]); cc.font=nf; cc.border=bd; cc.alignment=rgt; cc.number_format=WON
    rr+=1
rw=[6,6,9,34,24,12,11,11,11]+[9]*10
for j,w in enumerate(rw,1): raw.column_dimensions[get_column_letter(j)].width=w
raw.freeze_panes="D4"

# ===== 안내 =====
g=wb.create_sheet("안내_분류기준")
g.column_dimensions["A"].width=3; g.column_dimensions["B"].width=115
lines=[("title","현대로템 분기별 단일판매·공급계약 매출인식 — 작성 안내"),
 ("h","1. 데이터 범위와 출처"),
 ("p","· 출처: 현대로템 각 정기보고서의 「XI. 그 밖에 투자자 보호를 위하여 필요한 사항 → 1. 공시내용 진행 및 변경사항 → 나. 단일판매·공급계약 공시 관련 진행 사항 (1)」 표."),
 ("p","· 이 표는 2024 사업보고서(FY2024 실적, 2025.03 공시)부터 신설되었습니다. 그 이전 정기보고서(2021~2024.3Q)에는 이 표가 존재하지 않습니다(전수 확인함)."),
 ("p","· 표가 실제 존재하는 보고서 5개만 사용: 2024사업(FY24) · 2025반기(H1) · 2025.3Q · 2025사업(FY25) · 2026.1Q."),
 ("p","· 2025.1Q 분기보고서는 '본 항목은 반기·사업보고서에만 기재'로 표를 생략함 → 2025년 1분기 단독값은 존재하지 않아 '2025 상반기(1~6월)' 합산으로 제공합니다."),
 ("h","2. '분기별 매출인식' 산출 방법"),
 ("p","· 보고서의 '당기' 열 = 해당 회계연도 누계(YTD) 매출인식액. (검증: 싱가포르 LTA 2025.3Q당기 73,959 − 2025반기당기 43,019 = 30,940 ≈ 누적차분 30,939)"),
 ("p","· 달력분기값 = 당기(YTD) 차분:  2025.3Q = 당기(3Q) − 당기(반기),  2025.4Q = 당기(FY25) − 당기(3Q).  2024년=FY24 당기(연간), 2026.1Q=당기(3개월)."),
 ("p","· '당기'를 쓰는 이유: 누적 단순차분은 외화계약 환율 재환산·신규 표편입으로 왜곡되나 '당기'는 실제 손익 인식액이라 정확합니다."),
 ("p","· 분기값이 음수인 경우(예: 대만철도)는 원화 재환산·진행률 조정에 따른 것으로 보고서 원자료 그대로입니다."),
 ("h","3. 분류 기준"),
 ("p","· 대분류: 철도 / 방산 / 기타(제철설비·발전·자동차 생산설비).  소분류: 국내 / 국외(국외는 국가 표기)."),
 ("p","· 방산 국외 = 폴란드 K2전차(폴란드), 알타이전차 부품(터키 BMC).  방산 국내 = 방위사업청·한화향(대공포·장갑차·지휘소차·전차정비)."),
 ("p","· 미국 MBTA·LA메트로는 발주처가 '현대로템 미국법인'으로 표기되나 최종수요처 기준 미국(국외)으로 분류했습니다."),
 ("h","4. 검토 권장(메인 '비고' 열)"),
 ("p","· '통합계약': rowspan으로 금액이 하나로 묶여 공시된 건(호주 시드니 2층 전동차 공급/추가공급/개조, 현대제철 CDQ 내자/기전) — 중복 방지 위해 1건으로 집계."),
 ("p","· '표편입 2025.3Q': 그 시점에 처음 표에 등장 — 해당 분기값에 이전 기간분이 포함됐을 수 있어 확인 권장."),
 ("p","· '계약총액 외화기준': 2026.1Q부터 서식이 외화(천단위) 표기로 변경 — 원화 계약총액은 직전 사업보고서 기준으로 채웠습니다."),
 ("p","· 한국철도공사 'EMU-260 납품'은 신고일 2021-12-24 / 2024-07-29 의 서로 다른 2개 계약이라 2행으로 분리했습니다."),
 ("h","5. 시트 구성"),
 ("p","· [계약별_분기매출인식] 메인 표(분류·계약총액·최신누적·진행률·분기별 인식액, 대분류별 소계/총계)."),
 ("p","· [원자료_당기누적] 각 보고서에서 직접 추출한 당기(YTD)·누적 원본값(검증·재계산용)."),
 ("p","※ 모든 금액은 원본 보고서 표기 그대로이며 단위는 백만원입니다 (100백만원 = 1억원)."),]
i=1
for kind,txt in lines:
    cc=g.cell(i,2,txt)
    if kind=="title": cc.font=Font(name=FONT,bold=True,size=14,color="1F3864"); g.row_dimensions[i].height=24
    elif kind=="h": cc.font=Font(name=FONT,bold=True,size=11,color="2E5496"); g.row_dimensions[i].height=20
    else: cc.font=Font(name=FONT,size=10); cc.alignment=Alignment(wrap_text=True,vertical="top"); g.row_dimensions[i].height=15
    i+=1

out="/home/user/jooho96/현대로템_분기별_단일판매공급계약_매출인식.xlsx"
wb.save(out)
print("saved:", out, "| 계약수:", len(recs))
